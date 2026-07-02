from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

wb = Workbook()
ws = wb.active
ws.title = "UMass Contacts"

BRAND = "E2483A"; INK="171613"; CREAM="FBF7F0"; AMBER="F0A04B"; LINE="E7E0D5"
white = Font(name="Arial", color="FFFFFF", bold=True, size=11)
hdrfill = PatternFill("solid", fgColor=INK)
grpfill = PatternFill("solid", fgColor=BRAND)
p1fill = PatternFill("solid", fgColor="FCE9E7")
thin = Side(style="thin", color=LINE)
border = Border(left=thin,right=thin,top=thin,bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
wraptop = Alignment(wrap_text=True, vertical="top", horizontal="left")

cols = ["Priority","Name","Role","Org / Dept","Why they matter (your angle)","Best contact","Profile / Source","Verified?","Status (you fill)","Notes"]
widths = [9,20,30,22,40,30,42,11,16,34]
for i,w in enumerate(widths,1):
    ws.column_dimensions[chr(64+i)].width = w

def group(title):
    r = ws.max_row+1
    ws.cell(r,1,title)
    ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=len(cols))
    c=ws.cell(r,1); c.font=white; c.fill=grpfill; c.alignment=Alignment(vertical="center"); ws.row_dimensions[r].height=20

def hdr():
    r=ws.max_row+1
    for i,t in enumerate(cols,1):
        c=ws.cell(r,i,t); c.font=white; c.fill=hdrfill; c.alignment=Alignment(wrap_text=True,vertical="center"); c.border=border
    ws.row_dimensions[r].height=30

def row(vals, p1=False):
    r=ws.max_row+1
    for i,v in enumerate(vals,1):
        c=ws.cell(r,i,v); c.alignment=wraptop; c.border=border; c.font=Font(name="Arial",size=10)
        if p1: c.fill=p1fill
    ws.row_dimensions[r].height=46

# Title
ws.cell(1,1,"unscrewed.lol — UMass Amherst Launch Contacts")
ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(cols))
t=ws.cell(1,1); t.font=Font(name="Arial",bold=True,size=15,color=BRAND); t.alignment=Alignment(vertical="center"); ws.row_dimensions[1].height=26
ws.cell(2,1,"Compiled from public UMass pages, July 2026. Emails marked 'via profile' are Cloudflare-obfuscated on the site — click the profile link to reveal, or connect the Chrome extension and I'll resolve them.")
ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=len(cols))
ws.cell(2,1).font=Font(name="Arial",italic=True,size=9,color="57534E"); ws.cell(2,1).alignment=Alignment(wrap_text=True,vertical="center"); ws.row_dimensions[2].height=28

# GROUP 1
group("① SUSTAINABILITY / NEW2U  —  your #1 partner. Start here.")
hdr()
row(["1 ★","Laurie Simmons","Assistant Campus Sustainability Manager","Office of Sustainability","Her profile literally says 'Ask me about: New2U and Zero Waste.' The single best first email — owns the program day-to-day.","via profile page","https://www.umass.edu/sustainability/about/directory/laurie-simmons","No — obfuscated","","Send the New2U partnership email here first."], p1=True)
row(["1 ★","Ezra Small","Campus Sustainability Manager","Office of Sustainability","Oversees New2U + Zero Waste; the decision-maker. CC him or use as escalation.","'Email Ezra Small' link on FM Sustainability Team page","https://www.umass.edu/sustainability/our-commitment/facilities-management-sustainability-team","No — obfuscated","","CC on the Laurie email."], p1=True)
row(["1","Maitri Chandrashekar '26","New2U Marketing & Events Coordinator","Office of Sustainability (student)","Plans events in the store space and 'collaborates with RSOs, artists, makers to create community space' — perfect co-host for a Trade Day.","via profile page","https://www.umass.edu/sustainability/about/directory/maitri-chandrashekar-26","No — obfuscated","","Ideal partner for the flagship event."])
row(["2","Eva Bergloff '26","New2U Thrift Store Manager","Office of Sustainability (student)","On-the-ground store manager; student voice + move-out logistics.","via profile page","https://www.umass.edu/sustainability/about/directory/eva-bergloff-26","No — obfuscated","","Good ally / potential first ambassador."])
row(["3","Danush Aragonda","New2U Inventory & Pricing Coordinator","Office of Sustainability (student)","Runs the inventory platform — understands the item flow you'd plug into.","via profile page","https://www.umass.edu/sustainability/about/directory/danush-aragonda","No — obfuscated","","Secondary."])
row(["3","Marie Cloherty '27","Zero Waste Sustainability Fellow","Office of Sustainability (student)","Zero-waste focus; natural evangelist for the mission.","via profile page","https://www.umass.edu/sustainability/about/directory/marie-cloherty-27","No — obfuscated","","Secondary."])
row(["—","Office of Sustainability (general)","Dept. contact","360 Campus Center Way, Amherst MA 01003","Fallback if individual emails don't resolve; also socials to tag.","'Contact Us' link on site; IG @UMass_Sustain; FB /UMassSustainability","https://www.umass.edu/sustainability","Dept. channel","","Follow/DM their Instagram too."])

# GROUP 2
group("② RESIDENCE LIFE  —  dorm access + move-out programming (esp. Southwest)")
hdr()
row(["1","Jean A. MacKimmie","Director of Residential Life","Residential Life","Approves flyering in dorms and a 'dorm trade night' as RA programming. Southwest = densest launch zone.","jamackimmie@umass.edu · 413-545-6923","https://www.umass.edu/living/about/contact-us","Yes (from directory)","","Verify before sending; pitch dorm programming angle."], p1=True)
row(["1","Residential Life Main Office","Dept. contact","Berkshire House, 2nd Floor","Reliable front door; they route to RA/area coordinators.","living@umass.edu · 413-545-6923","https://www.umass.edu/living/about/contact-us","Yes","","Ask to be connected to a Southwest area coordinator."], p1=True)

# GROUP 3
group("③ PRESS — free, trusted reach once you have a trade or two to show")
hdr()
row(["2","Johnny Depin","Managing Editor","Massachusetts Daily Collegian","Managing editor address is the published contact for story pitches.","managingeditor@dailycollegian.com · 413-545-1809","https://dailycollegian.com/contact-us/","Yes","","Pitch the 'students trading, not trashing' story."])
row(["2","Caitlin Reardon","Editor-in-Chief","Massachusetts Daily Collegian","EIC; senior journalism major. Good named contact.","via Collegian staff page","https://dailycollegian.com/staff/","No","","Reference her by name in the pitch."])
row(["3","Collegian Newsroom","News desk","Campus Center Basement","Walk-in newsroom; sustainability/student-life beat.","413-545-1809","https://dailycollegian.com/staff/","Dept. channel","","In-person drop-by works on campus papers."])

# GROUP 4
group("④ STUDENT GOVERNMENT & CLUBS — to confirm (extension can pull these fast)")
hdr()
row(["2","SGA Sustainability / Student Life officer","(role — confirm current holder)","Student Government Association","Endorsement + possible small event funding + all-student reach.","To find: SGA site / RSO directory","https://www.umass.edu/sga/","No — TBD","","Have me pull the current officer via the extension."])
row(["2","Environmental / thrift / mutual-aid club leaders","Club presidents","Registered Student Orgs (RSOs)","Pre-assembled aligned audiences; make leaders your first ambassadors.","To find: Campus Pulse / SBS Pathways RSO directory","https://sbspathways.umass.edu/organizations/","No — TBD","","Prioritize sustainability & swap/thrift clubs."])

# Sheet 2: sequencing
ws2 = wb.create_sheet("How to use")
ws2.column_dimensions['A'].width=4; ws2.column_dimensions['B'].width=100
ws2.cell(1,1,""); ws2.cell(1,2,"How to work this list").font=Font(name="Arial",bold=True,size=14,color=BRAND)
steps=[
 "1. Send the New2U partnership email (in outreach/NEW2U_PARTNERSHIP_EMAIL.md) to Laurie Simmons, CC Ezra Small. This is your single highest-leverage action.",
 "2. Same week, email Residential Life (Jean MacKimmie / living@umass.edu) with the 'dorm trade night + Southwest flyering' angle.",
 "3. Once you have even one real trade to point to, pitch the Daily Collegian managing editor.",
 "4. Loop Maitri Chandrashekar in to co-host a Trade Day in/near the New2U store space.",
 "5. Ask me to pull the current SGA sustainability officer and the environmental/thrift RSO leaders (needs the Chrome extension connected).",
 "",
 "Email etiquette: short, one clean link (unscrewed.lol), lead with waste-reduction + student savings, never with ideology. One follow-up after ~5 business days if no reply.",
 "Emails marked 'via profile': open the profile link and click the email icon to reveal, OR connect the Chrome extension and I'll resolve and verify every one automatically.",
]
for i,s in enumerate(steps,3):
    c=ws2.cell(i,2,s); c.alignment=Alignment(wrap_text=True,vertical="top"); c.font=Font(name="Arial",size=10); ws2.row_dimensions[i].height=30

ws.freeze_panes="A3"
wb.save("UMass_Launch_Contacts.xlsx")
print("saved")
